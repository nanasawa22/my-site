# -*- coding: utf-8 -*-
"""Build the .NEXT On Tour Tokyo 2026 deck onto the supplied brand template."""
import copy, json, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from lxml import etree

SRC = 'base.pptx'
OUT = 'NEXT2026_AIGateway_draft.pptx'
ICONS = json.load(open('icons.json'))
ICON_DIR = 'icons'

CHARCO = RGBColor(0x13, 0x13, 0x13)
IRIS   = RGBColor(0xC6, 0xBE, 0xFE)
AQUA   = RGBColor(0xB2, 0xF8, 0xFF)
L1     = RGBColor(0xD7, 0xD7, 0xD7)
L3     = RGBColor(0xBF, 0xBF, 0xBF)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
MUTED  = RGBColor(0x4A, 0x4A, 0x55)

JP = 'Noto Sans JP'
EN = 'Montserrat'

# ── vertical rhythm (inches) ──────────────────────────────────────
TOP      = 1.40     # content top (no lead)
TOP_LEAD = 1.52     # content top when a lead line is present
BOT      = 4.42     # content bottom (above the conclusion strip)
BOT_SRC  = 4.16     # content bottom when a source line is present
LINE_Y   = 4.52     # conclusion rule
CONCL_Y  = 4.58
SRC_Y    = 4.88
LEFT     = 0.46
WIDTH    = 9.08

prs = Presentation(SRC)
LAYOUTS = {}
for m in prs.slide_masters:
    for l in m.slide_layouts:
        LAYOUTS.setdefault(l.name, l)

# ── helpers ───────────────────────────────────────────────────────
def set_font(run, size, bold=False, color=CHARCO, italic=False):
    f = run.font
    f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color
    f.name = EN
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        for e in rPr.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}' + tag.split(':')[1]):
            rPr.remove(e)
    ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    ea = etree.SubElement(rPr, '{%s}ea' % ns); ea.set('typeface', JP)
    cs = etree.SubElement(rPr, '{%s}cs' % ns); cs.set('typeface', JP)


def tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return s, tf


def para(tf, text, size, bold=False, color=CHARCO, first=False, space_after=0,
         bullet=False, italic=False, align=PP_ALIGN.LEFT, line=1.25):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    p.space_after = Pt(space_after)
    if bullet:
        text = '■  ' + text
    r = p.add_run(); r.text = text
    set_font(r, size, bold, color, italic)
    return p


def rich(tf, chunks, size, first=False, space_after=0, bullet=False, line=1.25):
    """chunks: list of (text, bold, hl) where hl in (None,'iris','aqua')"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.line_spacing = line
    p.space_after = Pt(space_after)
    if bullet:
        r = p.add_run(); r.text = '■  '; set_font(r, size, False, CHARCO)
    for text, bold, hl in chunks:
        r = p.add_run(); r.text = text
        set_font(r, size, bold, CHARCO)
        if hl:
            rPr = r._r.get_or_add_rPr()
            ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
            hi = etree.SubElement(rPr, '{%s}highlight' % ns)
            clr = etree.SubElement(hi, '{%s}srgbClr' % ns)
            clr.set('val', 'C6BEFE' if hl == 'iris' else 'B2F8FF')
    return p


def title(slide, text, lead=None):
    try:
        ph = slide.placeholders[0]
        ph.top, ph.left, ph.width = Inches(0.42 if not lead else 0.54), Inches(LEFT), Inches(WIDTH)
        ph.height = Inches(0.92)
        tf = ph.text_frame; tf.clear(); tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_bottom = 0
        para(tf, text, 28, True, CHARCO, first=True, line=1.15)
    except (KeyError, IndexError):
        _, tf = tb(slide, LEFT, 0.42 if not lead else 0.54, WIDTH, 0.92)
        para(tf, text, 28, True, CHARCO, first=True, line=1.15)
    if lead:
        _, tf2 = tb(slide, LEFT, 0.26, WIDTH, 0.24)
        para(tf2, lead, 15, False, MUTED, first=True)


def conclusion(slide, text):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    Inches(LEFT), Inches(LINE_Y),
                                    Inches(LEFT + WIDTH), Inches(LINE_Y))
    ln.line.color.rgb = CHARCO; ln.line.width = Pt(2)
    chip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(LEFT), Inches(CONCL_Y),
                                  Inches(0.52), Inches(0.24))
    chip.fill.solid(); chip.fill.fore_color.rgb = CHARCO; chip.line.fill.background()
    ctf = chip.text_frame; ctf.clear()
    ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
    ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(ctf, '結論', 10, True, WHITE, first=True, align=PP_ALIGN.CENTER)
    _, tf = tb(slide, LEFT + 0.64, CONCL_Y - 0.01, WIDTH - 0.64, 0.30)
    para(tf, text, 14, True, CHARCO, first=True)


def source(slide, text):
    _, tf = tb(slide, LEFT, SRC_Y, WIDTH, 0.30)
    para(tf, text, 10, False, MUTED, first=True, line=1.15)


def box(slide, x, y, w, h, fill=None, line=CHARCO, lw=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    return sh, tf


def arrow(slide, x1, y1, x2, y2, w=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = CHARCO; c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    tail = etree.SubElement(ln, '{%s}tailEnd' % ns)
    tail.set('type', 'triangle'); tail.set('w', 'med'); tail.set('len', 'med')
    return c


def icon(slide, slug, x, y, size=0.44):
    path = os.path.join(ICON_DIR, slug + '.png')
    if os.path.exists(path):
        slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(size), Inches(size))


def new(layout_name):
    return prs.slides.add_slide(LAYOUTS[layout_name])


def footer(slide):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx == 11:
            ph.text_frame.text = 'Nutanix Confidential'
            for p in ph.text_frame.paragraphs:
                for r in p.runs:
                    set_font(r, 9, False, RGBColor(0x6E, 0x6E, 0x78))


def strip_empty(slide):
    """remove untouched content placeholders so 'Click to add text' never shows"""
    for ph in list(slide.placeholders):
        idx = ph.placeholder_format.idx
        if idx in (11, 12):
            continue
        if not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


# ══════════════════════════════════════════════════════════════════
#  slide builders
# ══════════════════════════════════════════════════════════════════
IDC_FULL = ('Source: IDC, 2026 ｜ IDC Japan『2026年 国内AIインフラおよびAI向けITインフラ'
            'サービス市場動向分析』（Doc #JPJ54233926, March 2026）')

def s_bullets(t, lead, items, concl, src=None):
    sl = new('タイトルとコンテンツ'); title(sl, t, lead)
    top = TOP_LEAD if lead else TOP
    bottom = BOT_SRC if src else BOT
    ph = None
    for p in list(sl.placeholders):
        if p.placeholder_format.idx == 1: ph = p
    if ph is not None: ph._element.getparent().remove(ph._element)
    _, tf = tb(sl, LEFT, top, WIDTH, bottom - top)
    for i, chunks in enumerate(items):
        rich(tf, chunks, 24, first=(i == 0), space_after=9, bullet=True, line=1.2)
    conclusion(sl, concl)
    if src: source(sl, src)
    footer(sl); strip_empty(sl); return sl


def s_compare(t, lead, lh, litems, rh, ritems, concl, src=None):
    sl = new('Title and Comparison'); title(sl, t, lead)
    for p in list(sl.placeholders):
        if p.placeholder_format.idx in (1, 2, 13, 14):
            p._element.getparent().remove(p._element)
    top = TOP_LEAD if lead else TOP
    bottom = BOT_SRC if src else BOT
    cw = (WIDTH - 0.60) / 2
    for x, head, items, alt in ((LEFT, lh, litems, False), (LEFT + cw + 0.60, rh, ritems, True)):
        _, htf = tb(sl, x, top, cw, 0.28)
        para(htf, head, 14, True, CHARCO, first=True)
        ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(top + 0.30),
                                     Inches(x + cw), Inches(top + 0.30))
        ln.line.color.rgb = RGBColor(0x4A, 0x3F, 0xA8) if alt else CHARCO
        ln.line.width = Pt(2)
        _, btf = tb(sl, x, top + 0.42, cw, bottom - top - 0.42)
        for i, chunks in enumerate(items):
            rich(btf, chunks, 18, first=(i == 0), space_after=8, bullet=True, line=1.2)
    conclusion(sl, concl)
    if src: source(sl, src)
    footer(sl); strip_empty(sl); return sl


def s_grid(t, lead, cells, concl, cols=3, src=None, icons=None):
    """cells: list of (num|None, title, desc, fill)"""
    sl = new('Grid_Six Items' if cols == 3 else 'Title and Content 2x2'); title(sl, t, lead)
    for p in list(sl.placeholders):
        if p.placeholder_format.idx not in (0, 11, 12):
            p._element.getparent().remove(p._element)
    top = TOP_LEAD if lead else TOP
    bottom = BOT_SRC if src else BOT
    rows = (len(cells) + cols - 1) // cols
    gap = 0.20
    cw = (WIDTH - gap * (cols - 1)) / cols
    ch = (bottom - top - gap * (rows - 1)) / rows
    for i, (n, ct, cd, fill) in enumerate(cells):
        r, c = divmod(i, cols)
        x = LEFT + c * (cw + gap); y = top + r * (ch + gap)
        sh, tf = box(sl, x, y, cw, ch, fill=fill)
        ic_slug = (icons or {}).get(ct)
        pad = 0.0
        if ic_slug:
            icon(sl, ic_slug, x + 0.10, y + 0.08, 0.34)
            pad = 0.40
        tf.margin_top = Inches(0.05 + pad)
        f = True
        if n:
            para(tf, n, 10, True, RGBColor(0x5A, 0x5A, 0x66), first=True); f = False
        para(tf, ct, 15 if cols == 3 else 15, True, CHARCO, first=f, space_after=3, line=1.15)
        if cd:
            para(tf, cd, 11.5, False, RGBColor(0x2C, 0x2C, 0x36), line=1.25)
    conclusion(sl, concl)
    if src: source(sl, src)
    footer(sl); strip_empty(sl); return sl


def s_table(t, lead, rows, concl, src=None):
    sl = new('白紙'); title(sl, t, lead)
    top = TOP_LEAD if lead else TOP
    bottom = BOT_SRC if src else BOT
    n = len(rows)
    tbl_h = min(bottom - top, 0.42 * n)
    gt = sl.shapes.add_table(n, 2, Inches(LEFT), Inches(top), Inches(WIDTH), Inches(tbl_h)).table
    gt.columns[0].width = Inches(3.10); gt.columns[1].width = Inches(WIDTH - 3.10)
    for ri, (a, b, style) in enumerate(rows):
        for ci, txt in enumerate((a, b)):
            cell = gt.cell(ri, ci)
            cell.fill.background()
            cell.margin_left = Inches(0.04); cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.clear(); tf.word_wrap = True
            if style == 'head':
                para(tf, txt, 11, True, MUTED, first=True)
            elif style == 'road':
                para(tf, txt, 15, False, RGBColor(0x3A, 0x3A, 0x46), first=True, italic=(ci == 0))
            else:
                para(tf, txt, 15, ci == 0, CHARCO, first=True)
    conclusion(sl, concl)
    if src: source(sl, src)
    footer(sl); strip_empty(sl); return sl


def s_line_chart(t, lead, cats, series, concl, src, y=TOP, h=None):
    sl = new('白紙'); title(sl, t, lead)
    bottom = BOT_SRC if src else BOT
    cd = CategoryChartData(); cd.categories = cats
    for name, vals in series: cd.add_series(name, vals)
    gf = sl.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(LEFT), Inches(y),
                             Inches(WIDTH), Inches((h or (bottom - y))), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    for p in ch.plots:
        p.has_data_labels = False
    for i, s in enumerate(ch.series):
        s.format.line.color.rgb = CHARCO
        s.format.line.width = Pt(3 if i else 1.75)
        s.smooth = False
        if i == 0:
            ln = s.format.line._get_or_add_ln()
            ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
            d = etree.SubElement(ln, '{%s}prstDash' % ns); d.set('val', 'dash')
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = L1
    ch.value_axis.major_gridlines.format.line.width = Pt(0.75)
    ch.category_axis.has_major_gridlines = False
    for ax in (ch.value_axis, ch.category_axis):
        ax.format.line.color.rgb = CHARCO
        ax.tick_labels.font.size = Pt(10)
        ax.tick_labels.font.color.rgb = CHARCO
        ax.tick_labels.font.name = EN
    ch.font.size = Pt(10); ch.font.name = EN
    conclusion(sl, concl)
    if src: source(sl, src)
    footer(sl); strip_empty(sl); return sl


# ══════════════════════════════════════════════════════════════════
#  diagrams  (規程2.8 — shapes and text only, no icons)
# ══════════════════════════════════════════════════════════════════
def d_stack(t, lead, concl):
    """16: Nutanix Agentic AI stack"""
    sl = new('白紙'); title(sl, t, lead)
    sh, tf = box(sl, LEFT, 1.28, WIDTH, 1.42, fill=IRIS, lw=2)
    para(tf, 'Nutanix Enterprise AI (NAI) ― 中央AIコントロールプレーン', 12, True, CHARCO, first=True)
    for x, nm, l1, l2 in ((0.70, 'Agent Gateway', '統合エンドポイント／レートリミット', 'MCP接続制御／監査ログ'),
                          (5.06, 'Inference Management', 'モデル配備・GPUスケジューリング', 'オートスケール・ヘルスチェック')):
        _, itf = box(sl, x, 1.62, 3.94, 0.98, fill=WHITE)
        para(itf, nm, 13, True, CHARCO, first=True, align=PP_ALIGN.CENTER, space_after=3)
        para(itf, l1, 10, False, CHARCO, align=PP_ALIGN.CENTER, line=1.2)
        para(itf, l2, 10, False, CHARCO, align=PP_ALIGN.CENTER, line=1.2)
    arrow(sl, 5.00, 2.72, 5.00, 2.90, 2)
    _, k = box(sl, LEFT, 2.92, WIDTH, 0.62, fill=L1)
    para(k, 'Nutanix Kubernetes Platform (NKP) ／ CNCF準拠 Kubernetes', 12, True, CHARCO, first=True, space_after=2)
    para(k, 'Rancher・ベアメタル・AWS / Azure / GCP のコンテナ基盤でも可', 10, False, CHARCO)
    for x, nm, d in ((LEFT, 'Nutanix Cloud Platform', 'コンピュート・ネットワーク'),
                     (5.10, 'Nutanix Unified Storage (NUS)', 'モデル重み・データ')):
        _, b = box(sl, x, 3.62, 4.44, 0.60, fill=L3)
        para(b, nm, 12, True, CHARCO, first=True, space_after=2)
        para(b, d, 10, False, CHARCO)
    conclusion(sl, concl); footer(sl); strip_empty(sl); return sl


def d_lbfb(t, lead, concl):
    """20: LB vs fallback state transitions — no red/green"""
    sl = new('白紙'); title(sl, t, lead)
    panes = [('① 正常時 ― 分散', L1, 'normal'), ('② 劣化時 ― 切り離し', L1, 'degrade'), ('③ 障害時 ― 退避', IRIS, 'failover')]
    pw = (WIDTH - 0.36) / 3
    for i, (head, fill, mode) in enumerate(panes):
        x = LEFT + i * (pw + 0.18)
        box(sl, x, 1.28, pw, 2.86, fill=fill, lw=2 if mode == 'failover' else 1)
        _, htf = tb(sl, x + 0.12, 1.38, pw - 0.24, 0.24)
        para(htf, head, 12, True, CHARCO, first=True)
        _, g = box(sl, x + 0.14, 1.72, 1.06, 0.40, fill=WHITE)
        para(g, 'Gateway', 10, False, CHARCO, first=True, align=PP_ALIGN.CENTER)
        ys = (2.30, 2.92, 3.54)
        labels = ('A', 'B', 'C')
        for j, (yy, lb) in enumerate(zip(ys, labels)):
            if mode == 'degrade' and j == 0:
                _, bx = box(sl, x + pw - 0.86, yy, 0.62, 0.36, fill=L3)
                para(bx, '× ' + lb, 10, True, CHARCO, first=True, align=PP_ALIGN.CENTER)
                c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + 1.22), Inches(1.92),
                                            Inches(x + pw - 0.88), Inches(yy + 0.18))
                c.line.color.rgb = CHARCO; c.line.width = Pt(1.5)
                ln = c.line._get_or_add_ln()
                ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
                d = etree.SubElement(ln, '{%s}prstDash' % ns); d.set('val', 'dash')
                continue
            if mode == 'failover' and j == 0:
                _, bx = box(sl, x + pw - 0.86, yy, 0.62, 0.36, fill=L3)
                para(bx, '× ' + lb, 10, True, CHARCO, first=True, align=PP_ALIGN.CENTER)
                continue
            if mode == 'failover' and j == 2:
                _, bx = box(sl, x + pw - 0.86, yy, 0.62, 0.36, fill=WHITE)
                para(bx, lb, 10, False, CHARCO, first=True, align=PP_ALIGN.CENTER)
                continue
            fillc = AQUA if (mode == 'failover' and j == 1) else WHITE
            _, bx = box(sl, x + pw - 0.86, yy, 0.62, 0.36, fill=fillc,
                        lw=2 if fillc is AQUA else 1)
            para(bx, lb, 10, fillc is AQUA, CHARCO, first=True, align=PP_ALIGN.CENTER)
            arrow(sl, x + 1.22, 1.92, x + pw - 0.88, yy + 0.18)
    _, lg = tb(sl, LEFT, 4.20, WIDTH, 0.24)
    para(lg, '白＝稼働中　／　グレー＋×＝除外　／　実線＝トラフィックあり　／　破線＝停止　／　'
             '水色＝退避先（色に依存しない表現）', 10, False, MUTED, first=True)
    conclusion(sl, concl); footer(sl); strip_empty(sl); return sl


def d_ref(t, lead, concl):
    """28: reference architecture"""
    sl = new('白紙'); title(sl, t, lead)
    _, l = box(sl, LEFT, 1.72, 1.72, 1.44, fill=L1)
    para(l, '利用側', 11, True, CHARCO, first=True, space_after=4)
    for i, nm in enumerate(('業務アプリケーション', 'AIエージェント')):
        _, b = box(sl, LEFT + 0.14, 2.04 + i * 0.54, 1.44, 0.44, fill=WHITE)
        para(b, nm, 10, False, CHARCO, first=True, align=PP_ALIGN.CENTER)
    arrow(sl, 2.20, 2.44, 2.62, 2.44, 2)
    _, g = box(sl, 2.64, 1.50, 2.10, 1.90, fill=IRIS, lw=2)
    para(g, 'Agent Gateway', 13, True, CHARCO, first=True, align=PP_ALIGN.CENTER, space_after=5)
    for i, nm in enumerate(('統合エンドポイント', 'トークンレートリミット', 'MCP接続制御・監査')):
        _, b = box(sl, 2.78, 1.92 + i * 0.46, 1.82, 0.38, fill=WHITE)
        para(b, nm, 10, False, CHARCO, first=True, align=PP_ALIGN.CENTER)
    tgt = [('① 外部LLMプロバイダ', 'OpenAI / Anthropic / Azure / Bedrock ほか', 1.28),
           ('② NAI 上のセルフホストモデル', 'vLLM / NVIDIA NIM ― オンプレミス・エッジ', 2.22),
           ('③ MCPサーバー群', '業務ツール・内部データソース', 3.16)]
    for i, (nm, d, y) in enumerate(tgt):
        _, b = box(sl, 5.60, y, 3.94, 0.78, fill=AQUA)
        para(b, nm, 11, True, CHARCO, first=True, space_after=3)
        para(b, d, 10, False, CHARCO, line=1.2)
        arrow(sl, 4.74, 2.00 + i * 0.46, 5.58, y + 0.39, 2)
        cir = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.16), Inches(y + 0.22), Inches(0.26), Inches(0.26))
        cir.fill.solid(); cir.fill.fore_color.rgb = CHARCO; cir.line.fill.background()
        cir.shadow.inherit = False
        ctf = cir.text_frame; ctf.clear()
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(ctf, str(i + 1), 10, True, WHITE, first=True, align=PP_ALIGN.CENTER)
    conclusion(sl, concl); footer(sl); strip_empty(sl); return sl


def d_timeline(t, lead, steps, concl):
    """29: staged rollout"""
    sl = new('Timeline-2'); title(sl, t, lead)
    n = len(steps); gap = 0.18
    cw = (WIDTH - gap * (n - 1)) / n
    for i, (nm, d) in enumerate(steps):
        x = LEFT + i * (cw + gap)
        cir = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(1.90), Inches(0.34), Inches(0.34))
        cir.fill.solid(); cir.fill.fore_color.rgb = CHARCO; cir.line.fill.background()
        cir.shadow.inherit = False
        ctf = cir.text_frame; ctf.clear()
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(ctf, str(i + 1), 12, True, WHITE, first=True, align=PP_ALIGN.CENTER)
        ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(2.42),
                                     Inches(x + cw), Inches(2.42))
        ln.line.color.rgb = CHARCO; ln.line.width = Pt(2)
        _, tf = tb(sl, x, 2.52, cw, 1.30)
        para(tf, nm, 13, True, CHARCO, first=True, space_after=4)
        para(tf, d, 10, False, RGBColor(0x3A, 0x3A, 0x46), line=1.25)
    conclusion(sl, concl); footer(sl); strip_empty(sl); return sl


# ══════════════════════════════════════════════════════════════════
#  the deck
# ══════════════════════════════════════════════════════════════════
IDC_SURVEY_A = '「Japan AI Technology Strategy and Adoption Survey 2026」（n=1,250）を基に作成'
IDC_SURVEY_B = '「Japan Digital and AI Infrastructure Strategies and Investment Survey 2026」（n=551）を基に作成'


# ── 01 cover (fixed layout) ──
_c = new('表紙')
try:
    _t = _c.placeholders[0].text_frame; _t.clear(); _t.word_wrap = True
    para(_t, '「AIが止まれば業務が止まる」時代へ\n― 本番業務で使えるエンタープライズAIのアーキテクチャ',
         26, True, WHITE, first=True, line=1.2)
except Exception:
    pass
try:
    _s = _c.placeholders[1].text_frame; _s.clear(); _s.word_wrap = True
    para(_s, '部署名\n肩書き\n氏名', 13, False, RGBColor(0xC9, 0xC8, 0xD2), first=True, line=1.4)
except Exception:
    pass

# ── 02 speaker (fixed layout) ──
_sp = new('スピーカータイトル')
for _ph in list(_sp.placeholders):
    _i = _ph.placeholder_format.idx
    if _i == 2:
        _tf = _ph.text_frame; _tf.clear(); para(_tf, 'Name', 22, True, WHITE, first=True)
    elif _i == 3:
        _tf = _ph.text_frame; _tf.clear(); para(_tf, 'Dept, title', 12, False, RGBColor(0xC9,0xC8,0xD2), first=True)
    elif _i == 1:
        _tf = _ph.text_frame; _tf.clear(); para(_tf, 'Nutanix', 12, False, RGBColor(0xC9,0xC8,0xD2), first=True)


MAP_ROWS = [('本番AI基盤の非機能要件', '対応する NAI 2.7 の機能', 'head'),
            ('移植性', '統合エンドポイント ― 単一API・マルチプロバイダ', ''),
            ('可用性', '統合エンドポイント ― ロードバランス／フォールバック', ''),
            ('コスト予測可能性', 'トークンレートリミット', ''),
            ('データ主権', 'MCP接続制御', ''),
            ('（横断する最適化）', 'セマンティックルーティング【ロードマップ】', 'road')]

# 03
s_table('要件は4つ。充足する層はアプリではなく Gateway', '本日の結論', MAP_ROWS,
        'この対応関係が、本日の地図になる')
# 04
s_line_chart('2027年、推論が学習を上回る', '補助ツールから、業務プロセスの構成要素へ',
             ['2024', '2025', '2026', '2027', '2028', '2029', '2030'],
             [('学習向け', [305.6, 380.8, 411.2, 409.5, 405.3, 393.4, 385.9]),
              ('推論向け', [187.8, 292.7, 384.6, 418.4, 459.8, 500.5, 552.8])],
             'AI基盤は「作る」から「使い続ける」フェーズへ移った',
             'Source: IDC, 2026 ｜ ' + IDC_FULL.split('｜ ')[1] + ' の Table 3（単位：十億円）を基に作成',
             y=1.24, h=2.86)
# 05
s_compare('約束される水準と、実際に起きたこと', None,
          '提供されるSLA',
          [[('OpenAI Scale Tier ／ 99.9%', False, None)],
           [('Anthropic Priority Tier ／ 99.5%目標', False, None)],
           [('Azure OpenAI ／ 99.9%', False, None)],
           [('99.9% ＝ 年間 約8.76時間', True, 'iris')]],
          '実際の障害 ／ Azure OpenAI',
          [[('2026年5月29日 09:39–17:05 UTC', False, None)],
           [('1回で約7時間26分', True, 'iris')],
           [('レイテンシ増大・タイムアウト・5XX', False, None)],
           [('引き金は内部リトライの急増', True, 'aqua')]],
          '単一プロバイダのSLAは、基幹業務のSLOを満たさない',
          '出所：Microsoft Azure「Azure status history」／ OpenAI Help Center、各社公開ドキュメント。'
          '年間停止時間は 8,760時間 × 0.1% として算出（自社計算）')
# 06
s_grid('可用性の毀損は、サービス停止に限らない', None,
       [('01', 'サービス停止', 'プロバイダ側の障害', None),
        ('02', 'レート制限', '429。上限への到達', None),
        ('03', 'レイテンシ劣化', '応答するが間に合わない', None),
        ('04', '提供終了・版数変更', 'モデルの廃止と更新', None),
        ('05', 'リージョン障害', '地域単位での影響', None),
        (None, 'アプリから見れば\n5類型すべてが\n等価な「処理不能」', None, IRIS)],
       'アプリから見れば、5類型はすべて等価な「処理不能」である', cols=3)
# 07
s_grid('本番AI基盤に求められる4つの非機能要件', None,
       [(None, '可用性  Availability', '個別プロバイダの障害時にも処理を継続できること', None),
        (None, 'コスト予測可能性  Cost Predictability', '消費量を計測・按分・制限し、支出を統制下に置けること', None),
        (None, 'データ主権  Data Sovereignty', '機密データの所在と準拠法域を統制できること', None),
        (None, '移植性  Portability', '特定プロバイダに依存せず代替可能に保てること', None)],
       'この4つは、国内のAIインフラ意思決定者が実際に求めているものである', cols=2,
       src='Source: IDC, 2026 ｜ IDC Japan (Doc #JPJ54233926, March 2026)、' + IDC_SURVEY_B,
       icons={'可用性  Availability': 'availability',
              'コスト予測可能性  Cost Predictability': 'cost',
              'データ主権  Data Sovereignty': 'lock',
              '移植性  Portability': 'multicloud'})
# 08
s_compare('不確実性は二重である', 'コスト予測可能性',
          '① 量 ― 実行時に決まる',
          [[('推論 → 行動 → 観測 → 再推論', False, None)],
           [('反復回数はタスクの難易度に依存', False, None)],
           [('同一業務でも消費量が一定しない', True, 'iris')]],
          '② 価格 ― 為替で変動する',
          [[('海外プロバイダの課金はドル建て', False, None)],
           [('為替変動が利用料の急騰を招く', True, 'aqua')],
           [('ソブリンAI推進目的の第2位', False, None)]],
          '量と価格の双方が動くため、事前見積りは原理的に成立しない',
          'Source: IDC, 2026 ｜ IDC Japan (Doc #JPJ54233926, March 2026)、' + IDC_SURVEY_A + '。'
          '産業分野別では金融・流通／サービスが為替リスクを特に重視')
# 09
s_bullets('統制対象は、モデルだけではない', 'データ主権',
          [[('改正個人情報保護法、金融・医療の業法、GDPR', False, None)],
           [('機密データを外部LLMへ送出できない業務が実在する', False, None)],
           [('エージェントは業務ツール・内部DBにも接続する', True, 'iris')],
           [('モデルへの経路だけを統制しても、統制は完成しない', False, None)]],
          '統制すべき経路は、モデルへの経路だけではない',
          'Source: IDC, 2026 ｜ IDC Japan (Doc #JPJ54233926, March 2026)。'
          'ソブリンAI推進目的の第1位は「データ主権の確保」。' + IDC_SURVEY_A)
# 10
s_grid('信頼性設計の定石 ― 継続時間で担当が変わる', None,
       [('〜数秒', 'リトライ', '一過性のノイズ。瞬断、瞬間的な混雑を吸収する', None),
        ('数十秒〜数分', 'サーキットブレーカー', '劣化したエンドポイントを切り離し、連鎖を止める', None),
        ('数分〜数時間', 'フォールバック', '長期障害。別の供給元へ退避する', None)],
       '3つの機構は競合ではなく、障害の継続時間による分担である', cols=3,
       icons={'リトライ': 'flow', 'サーキットブレーカー': 'risk', 'フォールバック': 'availability'})

# 11 — the linchpin: chart on the right half
sl11 = new('Title and Comparison')
title(sl11, 'サーキットブレーカーは「障害」にしか反応しない')
for p in list(sl11.placeholders):
    if p.placeholder_format.idx in (1, 2, 13, 14):
        p._element.getparent().remove(p._element)
_cw = (WIDTH - 0.60) / 2
_, _h = tb(sl11, LEFT, TOP, _cw, 0.28); para(_h, '従来のCBが検知するもの', 14, True, CHARCO, first=True)
_l = sl11.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(LEFT), Inches(TOP + 0.30), Inches(LEFT + _cw), Inches(TOP + 0.30))
_l.line.color.rgb = CHARCO; _l.line.width = Pt(2)
_, _b = tb(sl11, LEFT, TOP + 0.44, _cw, 1.9)
for i, txt in enumerate(('エラー率の上昇', 'レイテンシの悪化', 'タイムアウトの頻発')):
    rich(_b, [(txt, False, None)], 18, first=(i == 0), space_after=9, bullet=True)
_x2 = LEFT + _cw + 0.60
_, _h2 = tb(sl11, _x2, TOP, _cw, 0.28); para(_h2, 'エージェントの暴走', 14, True, CHARCO, first=True)
_l2 = sl11.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(_x2), Inches(TOP + 0.30), Inches(_x2 + _cw), Inches(TOP + 0.30))
_l2.line.color.rgb = RGBColor(0x4A, 0x3F, 0xA8); _l2.line.width = Pt(2)
_cd = CategoryChartData(); _cd.categories = ['', '', '', '', '']
_cd.add_series('エラー率', [4, 5, 4, 5, 4]); _cd.add_series('トークン消費量', [5, 18, 44, 76, 100])
_gf = sl11.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(_x2), Inches(TOP + 0.40), Inches(_cw), Inches(1.72), _cd)
_c = _gf.chart; _c.has_title = False
_c.has_legend = True; _c.legend.position = XL_LEGEND_POSITION.BOTTOM; _c.legend.include_in_layout = False
_c.font.size = Pt(9); _c.font.name = EN
for i, s in enumerate(_c.series):
    s.format.line.color.rgb = CHARCO; s.smooth = False
    s.format.line.width = Pt(1.5 if i == 0 else 2.75)
    if i == 0:
        _ln = s.format.line._get_or_add_ln()
        _d = etree.SubElement(_ln, '{http://schemas.openxmlformats.org/drawingml/2006/main}prstDash')
        _d.set('val', 'dash')
_c.value_axis.has_major_gridlines = False
_c.value_axis.visible = False
_c.category_axis.has_major_gridlines = False
_, _n = tb(sl11, _x2, TOP + 2.24, _cw, 0.34)
rich(_n, [('200 OK のままCBは作動しない', True, 'iris')], 18, first=True, bullet=True)
conclusion(sl11, 'コストに対する遮断器が、既存の設計には存在しない')
footer(sl11); strip_empty(sl11)

# 12
s_grid('個別実装の限界と、その帰結', None,
       [(None, '組合せ爆発', 'N アプリ × M プロバイダ', None),
        (None, 'SDK差異の吸収', 'プロバイダごとの作法を各所で吸収', None),
        (None, '鍵の分散', 'APIキーがアプリごとに散在する', None),
        (None, '計測の分断', '全社のトークン使用量が可視化できない', None),
        (None, '変更の重さ', 'ポリシー変更に全アプリの再デプロイ', None),
        (None, '同じ仕組みを\nN回実装し、N回維持する', None, IRIS)],
       '同じ仕組みをN回実装しN回維持する構造は、いずれ破綻する', cols=3)
# 13
s_compare('既視感のある話 ― API Gateway が通った道', None,
          '10年前 ／ REST API',
          [[('生の接続性が先に普及した', False, None)],
           [('認証・流量制御・監査が各アプリに散在', False, None)],
           [('API Gateway が責務を引き受けた', True, 'iris')]],
          '現在 ／ エージェントとモデル',
          [[('生の接続性が先に普及している', False, None)],
           [('認証・流量制御・監査が各アプリに散在', False, None)],
           [('同じ構造が再現されている', True, 'aqua')]],
          'これは新しい問題ではなく、10年前に解いた問題の再来である',
          '参考：Forbes / Janakiram MSV「Agent Gateways Are Becoming The Control Plane For Enterprise AI」2026年7月5日')
# 14
s_grid('Gateway が引き受ける責務', None,
       [(None, '認証・認可', '誰がどのモデルを呼べるか', None),
        (None, 'ルーティング／ロードバランス', '平常時の分散', AQUA),
        (None, 'フォールバック', '異常時の退避', AQUA),
        (None, 'レート制限', '消費量の上限強制', AQUA),
        (None, '計測・監査ログ', '誰が何トークン使ったか', None),
        (None, 'MCP接続制御', 'ツールへの経路の統制', AQUA)],
       '信頼性・統制・可観測性は、アプリではなく基盤層の責務である', cols=3,
       icons={'認証・認可': 'auth', 'ルーティング／ロードバランス': 'network',
              'フォールバック': 'availability', 'レート制限': 'budget',
              '計測・監査ログ': 'measure', 'MCP接続制御': 'lock'})
# 15
s_bullets('エンタープライズ要件では、まだ足りない', None,
          [[('セルフホストモデルと外部モデルを', False, None), ('同一の統制下', True, 'iris'), ('に置けるか', False, None)],
           [('オンプレミス／エアギャップ環境で動作するか', False, None)],
           [('マルチテナントと部門別のコスト按分に対応できるか', False, None)],
           [('監査要件を満たすログを出せるか', False, None)]],
          'セルフホストと外部モデルを同一の統制下に置けるかが分岐点になる')
# 16
d_stack('Nutanix Agentic AI の全体像', 'NAI 2.7 ― AI Gateway から改名のうえ GA',
        'NAIが、インフラとエージェントの間に立つ中央コントロールプレーンとなる')
# 17
s_bullets('単一APIで、すべてのモデルへ', '統合エンドポイント 1／3',
          [[('Azure上のGPT-4、AnthropicのClaude、NAI上のLlama', False, None)],
           [('― すべて同一のエンドポイントを通る', True, 'iris')],
           [('Bedrock／NIM／Hugging Face／独自モデル。', False, None), ('検証済み74種', True, 'aqua')]],
          '多くの組織は両方を使う。ならば、束ねる層が要る',
          '配備形態の利用予定：パブリッククラウド 50%超／専有型ITインフラ 35〜45%（複数回答）。'
          'Source: IDC, 2026 ｜ IDC Japan (Doc #JPJ54233926, March 2026)、' + IDC_SURVEY_B)
# 18
s_bullets('ロードバランス ― 平常時の分散', '統合エンドポイント 2／3',
          [[('複数クラスタ・複数プロバイダに跨って負荷を分散する', True, 'iris')],
           [('リモートのNAIクラスタもプロバイダとして登録できる', True, 'aqua')],
           [('分散したGPUを活用し、ボトルネックを解消する', False, None)]],
          '平常時のスループットは、分散によって確保する',
          '※ 分散アルゴリズム・ヘルスチェック方式は公式ドキュメントで要確認')
# 19
s_bullets('フォールバック ― 異常時の退避', '統合エンドポイント 3／3',
          [[('複数の上流プロバイダをあらかじめ設定する', False, None)],
           [('プライマリのモデルエンドポイントが', False, None),
            ('障害を起こした場合、またはレート制限に到達した場合', True, 'iris')],
           [('健全なバックアップのモデルエンドポイントへ自動的にルーティング', False, None)],
           [('アプリケーションコードの改修は不要', True, 'aqua')]],
          '異常時の継続は、退避によって確保する。アプリの改修は不要')
# 20
d_lbfb('ロードバランスとフォールバックの関係', None,
       'LBは平常時の分散、フォールバックは異常時の退避。役割が異なる')

# 21 — the peak
s_compare('トークンレートリミット ＝ トークンのサーキットブレーカ', None,
          '従来のサーキットブレーカ',
          [[('トリガーは', False, None), ('エラー率', True, None), ('とレイテンシ', False, None)],
           [('守るのは', False, None), ('システムの健全性', True, None)],
           [('暴走時は作動しない', False, None)]],
          'トークンのサーキットブレーカ',
          [[('監視するのは', False, None), ('トークン消費量', True, 'iris')],
           [('守るのは', False, None), ('支出の予測可能性', True, 'iris')],
           [('200 OK のまま遮断できる', False, None)]],
          'エラー率ではなくトークン消費量を監視する遮断器が要る',
          'スライド11で示した「コストに対する遮断器が存在しない」への回答')
# 22
s_bullets('中央で強制し、単位ごとに可視化する', 'トークンレートリミット 2／2',
          [[('エンドポイント単位', True, 'iris'), ('と', False, None), ('ユーザー単位', True, 'iris'),
            ('でトークンベースのレート制限', False, None)],
           [('コストと流量のガバナンスを中央で強制する', False, None)],
           [('モデルベンダーを横断してトークン使用量を可視化', False, None)],
           [('目的は明快 ― ', False, None), ('bill shock を防ぐ', True, 'aqua')]],
          '見えない支出は制御できない。まず計測を中央に集める')
# 23
s_grid('トークン遮断の設計論点', None,
       [('WHO', '遮断の単位', 'ユーザー／APIキー／エージェント／チーム／時間窓', None),
        ('WHAT', '上限到達時の挙動', '429を返す／キューイングする／下位モデルへ降格する', IRIS),
        ('HOW MUCH', '閾値の設計', 'ハードリミットか、警告閾値を別に設けるか', None)],
       '遮断か降格かは、技術ではなく業務要件が決める', cols=3,
       icons={'遮断の単位': 'policies', '上限到達時の挙動': 'alerts', '閾値の設計': 'measure'})
# 24
s_bullets('MCP接続制御　[ Tech Preview ]', None,
          [[('APIキーごとに、アクセスできるツールを定義する', True, 'iris')],
           [('Tool-Level Filtering ― ツール単位で Read Only / Write を制御', False, None)],
           [('個々のMCPサーバーではなく、', False, None), ('ゲートウェイ側でAPIキーを注入', True, 'aqua')]],
          'エージェントに認証情報を持たせず、ツール単位で権限を絞る',
          '※ MCPサーバーガバナンスおよび同梱テストエージェントは Tech Preview であり、本番利用は想定されていません')
# 25
s_table('答え合わせ ― 4要件はすべて基盤層で充足できる', None,
        [('非機能要件', '対応する機能', 'head'),
         ('移植性', '統合エンドポイント ― 単一API・マルチプロバイダ', ''),
         ('可用性', '統合エンドポイント ― ロードバランス／フォールバック', ''),
         ('コスト予測可能性', 'トークンレートリミット（per endpoint / per user）', ''),
         ('データ主権', 'MCP接続制御（Tech Preview）', '')],
        '4要件はすべて、アプリを改修せず基盤層で充足できる',
        '実行環境：任意のCNCF準拠Kubernetes ／ オンプレミス・エッジ・NEO Cloud・パブリッククラウド')
# 26
s_bullets('セマンティックルーティング', None,
          [[('ここまでは、呼ぶモデルを', False, None), ('アプリが指定する', True, None), ('前提だった', False, None)],
           [('リクエストの', False, None), ('意味内容・複雑性・意図', True, 'aqua'), ('に基づき動的に選択する', False, None)],
           [('単純な要求は小型モデルへ、困難な要求は高性能モデルへ', False, None)]],
          'モデル選択をアプリから基盤へ移すことが、次の一手になる',
          '※ 本機能はロードマップです。提供時期・仕様は未確定です')
# 27
s_bullets('何が変わるか', None,
          [[('コスト・品質・レイテンシを', False, None), ('同時に最適化', True, 'iris'), ('できる', False, None)],
           [('上位モデルと小型モデルには大きな単価差がある', False, None)],
           [('相当割合の要求を小型モデルへ振り分けても品質を維持しうる', False, None)],
           [('ルーティングポリシーは一度定義すればよい', False, None)]],
          'コスト・品質・レイテンシは、トレードオフではなく同時に最適化できる',
          '※ 本機能はロードマップです。提供時期・仕様は未確定です')
# 28
d_ref('リファレンス構成', None, 'アプリが知るべきエンドポイントは、最後まで1つである')
# 29
d_timeline('段階的な導入手順', None,
           [('集約', '全AIトラフィックを統合エンドポイント経由に。可観測性の確保のみ'),
            ('制限', 'トークンレートリミットを適用する'),
            ('冗長化', 'ロードバランスとフォールバックを構成する'),
            ('内製化', 'セルフホストモデルを併用する'),
            ('統制', 'MCP接続制御を適用する')],
           'まず可視化だけを行う。それが最も確実な第一歩になる')
# 30
s_bullets('まとめ', None,
          [[('本番AI基盤の要件は', False, None), ('可用性・コスト予測可能性・データ主権・移植性', True, None)],
           [('充足すべき層は', False, None), ('アプリケーションではなく Gateway', True, 'iris')],
           [('次の一手はセマンティックルーティング（ロードマップ）', False, None)]],
          '止まらないことは、努力目標ではなく設計で満たす前提条件である')



# ── 31 thank you (fixed layout) ──
_th = new('1_Built for the Future_Thank You_Q&A')


prs.save(OUT)
print('saved %s — %d slides' % (OUT, len(prs.slides._sldIdLst)))
